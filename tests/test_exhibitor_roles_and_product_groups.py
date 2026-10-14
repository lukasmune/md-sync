from openpyxl import load_workbook

import database
import scraper
from excel_export import export_event
from product_groups import GIFA_PRODUCT_GROUP_COLUMNS, GIFA_PRODUCT_GROUPS


def setup_database(tmp_path, monkeypatch):
    monkeypatch.setattr(database, "DATABASE_PATH", tmp_path / "test.db")
    monkeypatch.setattr(database, "DATA_DIR", tmp_path)
    database.initialize_database()
    return database.get_or_create_exhibition(
        "GIFA", 2023, "GMTN2023.gifa"
    )


def test_source_co_exhibitors_are_expanded_with_parent_relationship():
    records = scraper.expand_exhibitor_records(
        {
            "exh": "DIRECT-1",
            "name": "Direct Company",
            "coExhibitors": [
                {"exh": "CO-1", "name": "Co Company"},
            ],
        }
    )

    assert [record["exhibitor_role"] for record in records] == [
        "direct",
        "co-exhibitor",
    ]
    assert records[1]["parent_source_exhibitor_id"] == "DIRECT-1"


def test_co_exhibitor_persists_without_merging_and_sales_fields_survive(
    tmp_path, monkeypatch
):
    exhibition_id = setup_database(tmp_path, monkeypatch)
    direct = {
        "exh": "DIRECT-1",
        "name": "Direct Company",
        "country": "Germany",
        "city": "Düsseldorf",
    }
    co = {
        "exh": "CO-1",
        "name": "Direct Company",
        "country": "Germany",
        "city": "Düsseldorf",
        "exhibitor_role": "co-exhibitor",
        "parent_source_exhibitor_id": "DIRECT-1",
    }
    first = database.sync_exhibitor(direct, exhibition_id)
    database.sync_exhibitor(co, exhibition_id)

    with database.get_connection() as connection:
        connection.execute(
            "UPDATE companies SET sales_status = 'Contacted' WHERE id = ?",
            (first["company_id"],),
        )
    database.sync_exhibitor(
        dict(co, name="Direct Company Updated"), exhibition_id
    )
    with database.get_connection() as connection:
        rows = connection.execute(
            """
            SELECT exhibitor_role, parent_source_exhibitor_id
            FROM exhibition_companies ORDER BY source_exhibitor_id
            """
        ).fetchall()
        sales_status = connection.execute(
            "SELECT sales_status FROM companies WHERE id = ?",
            (first["company_id"],),
        ).fetchone()["sales_status"]

    assert len(rows) == 2
    co_row = next(row for row in rows if row["exhibitor_role"] == "co-exhibitor")
    assert co_row["parent_source_exhibitor_id"] == "DIRECT-1"
    assert sales_status == "Contacted"


def test_duplicate_candidates_are_diagnostic_only(tmp_path, monkeypatch):
    exhibition_id = setup_database(tmp_path, monkeypatch)
    database.sync_exhibitor(
        {
            "exh": "A-1",
            "name": "Acme Foundry GmbH",
            "country": "Germany",
            "city": "Essen",
        },
        exhibition_id,
    )
    candidates = database.find_duplicate_candidates(
        "Acme Foundry GmbH", country="Germany", city="Essen"
    )

    assert len(candidates) == 1
    with database.get_connection() as connection:
        assert connection.execute(
            "SELECT COUNT(*) FROM companies"
        ).fetchone()[0] == 1


def test_gifa_product_groups_are_ordered_and_export_as_x_or_blank(
    tmp_path, monkeypatch
):
    exhibition_id = setup_database(tmp_path, monkeypatch)
    database.sync_exhibitor(
        {
            "exh": "GIFA-1",
            "name": "Categorized Foundry",
            "categories": [
                {
                    "id": "cat-1",
                    "label": "Foundry",
                    "hierarchy": [
                        {"id": "group-1", "label": GIFA_PRODUCT_GROUPS[0]},
                        {"id": "group-2", "label": GIFA_PRODUCT_GROUPS[4]},
                    ],
                }
            ],
        },
        exhibition_id,
    )
    path = export_event("GIFA", 2023, tmp_path)
    workbook = load_workbook(path, read_only=True)
    headers = [cell.value for cell in next(workbook["Sales Leads"].iter_rows())]
    row = dict(zip(
        headers,
        (cell.value for cell in next(
            workbook["Sales Leads"].iter_rows(min_row=2)
        )),
    ))

    assert headers[-29:] == list(GIFA_PRODUCT_GROUP_COLUMNS)
    assert row[GIFA_PRODUCT_GROUP_COLUMNS[0]] == "x"
    assert row[GIFA_PRODUCT_GROUP_COLUMNS[4]] == "x"
    assert row[GIFA_PRODUCT_GROUP_COLUMNS[1]] in ("", None)
    assert "Parent Exhibitor ID" not in headers
    assert GIFA_PRODUCT_GROUP_COLUMNS[0].startswith("1. ")
    assert GIFA_PRODUCT_GROUP_COLUMNS[-1].startswith("29. ")
    workbook.close()
