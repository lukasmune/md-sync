import pytest
from openpyxl import load_workbook

import database
import excel_export
import scraper


def _record(exh="E1", name="Acme GmbH", **values):
    record = {
        "type": "profile",
        "exh": exh,
        "exhSeoId": "SEO-" + exh,
        "name": name,
        "country": "Germany",
        "city": "Berlin",
        "location": "Hall 1 / A1",
        "premium": False,
        "tags": [],
        "eventIcons": [{"label": "GIFA 2023"}],
    }
    record.update(values)
    return record


@pytest.fixture
def report_context(tmp_path, monkeypatch):
    database.DATABASE_PATH = tmp_path / "sync.db"
    database.DATA_DIR = tmp_path
    reports = tmp_path / "changes"
    monkeypatch.setattr(
        scraper,
        "export_change_report",
        lambda run_id, changes, **kwargs: excel_export.export_change_report(
            run_id, changes, reports, **kwargs
        ),
    )
    return reports


def test_report_contains_new_company_and_participation(report_context):
    scraper.sync_event("GIFA", 2023, "GIFA2023", [_record()])
    report = next((report_context / "GIFA").glob("sync_changes_*.xlsx"))
    workbook = load_workbook(report, read_only=True)

    assert workbook["Summary"]["B2"].value == 1
    assert workbook["Summary"]["B3"].value == 1
    assert workbook["New Companies"].max_row == 2
    assert workbook["New Participations"].max_row == 2
    participation_headers = [
        cell.value for cell in workbook["New Participations"][1]
    ]
    assert "Parent Exhibitor ID" not in participation_headers
    assert "Tags" not in participation_headers


def test_known_company_in_new_edition_is_new_participation(report_context):
    scraper.sync_event("GIFA", 2023, "GIFA2023", [_record()])
    scraper.sync_event(
        "GIFA", 2027, "GIFA2027", [_record(exh="E2")]
    )
    reports = sorted((report_context / "GIFA").glob("sync_changes_*.xlsx"))
    workbook = load_workbook(reports[-1], read_only=True)

    assert workbook["Summary"]["B2"].value == 0
    assert workbook["Summary"]["B3"].value == 1
    assert workbook["New Participations"]["B2"].value == "GIFA"
    assert workbook["New Participations"]["C2"].value == 2027


def test_batch_sync_creates_one_report_per_affected_edition(
    report_context, monkeypatch, capsys
):
    records = {
        "GIFA 2023": [_record(exh="GIFA-1")],
        "METEC 2023": [_record(exh="METEC-1", name="Metal GmbH")],
    }
    monkeypatch.setattr(
        scraper,
        "fetch_exhibitors_by_event",
        lambda label, event_config=None: records[label],
    )

    scraper.sync_configured_events(["GIFA", "METEC"])
    output = capsys.readouterr().out

    gifa_reports = list((report_context / "GIFA").glob("sync_changes_*.xlsx"))
    metec_reports = list((report_context / "METEC").glob("sync_changes_*.xlsx"))
    assert len(gifa_reports) == 1
    assert len(metec_reports) == 1
    assert "GIFA 2023: New companies: 1 | New participations: 1 | Updated records: 0 | Changed fields: 0" in output
    assert "METEC 2023: New companies: 1 | New participations: 1 | Updated records: 0 | Changed fields: 0" in output
    assert "Change report:" in output

    for report in (gifa_reports + metec_reports):
        workbook = load_workbook(report, read_only=True)
        assert workbook["Summary"]["B2"].value == 1
        assert workbook["Summary"]["B3"].value == 1


def test_batch_sync_prints_no_changes_for_unchanged_editions(
    report_context, monkeypatch, capsys
):
    records = {"GIFA 2023": [_record(exh="GIFA-1")]}
    monkeypatch.setattr(
        scraper,
        "fetch_exhibitors_by_event",
        lambda label, event_config=None: records[label],
    )
    scraper.sync_configured_events(["GIFA"])
    capsys.readouterr()
    scraper.sync_configured_events(["GIFA"])

    assert capsys.readouterr().out.strip() == (
        "Sync completed. No changes detected."
    )
    assert len(list((report_context / "GIFA").glob("sync_changes_*.xlsx"))) == 1


def test_consecutive_changed_runs_keep_unique_report_files(report_context):
    scraper.sync_event(
        "GIFA", 2023, "GIFA2023", [_record(website="https://one.example")]
    )
    scraper.sync_event(
        "GIFA", 2023, "GIFA2023", [_record(website="https://two.example")]
    )
    scraper.sync_event(
        "GIFA", 2023, "GIFA2023",
        [_record(website="https://three.example")],
    )

    reports = list((report_context / "GIFA").glob("sync_changes_*.xlsx"))
    assert len(reports) == 3
    assert len({report.name for report in reports}) == 3
    assert all(
        report.name.startswith("sync_changes_GIFA_")
        and len(report.stem.rsplit("_", 1)[-1]) == 6
        for report in reports
    )


def test_multiple_changes_in_one_edition_are_consolidated(report_context):
    scraper.sync_event(
        "GIFA",
        2023,
        "GIFA2023",
        [
            _record(exh="E1", name="Acme GmbH"),
            _record(exh="E2", name="Beta GmbH"),
        ],
    )

    reports = list((report_context / "GIFA").glob("sync_changes_*.xlsx"))
    assert len(reports) == 1
    workbook = load_workbook(reports[0], read_only=True)
    assert workbook["Summary"]["B2"].value == 2
    assert workbook["Summary"]["B3"].value == 2
    assert workbook["New Companies"].max_row == 3
    assert workbook["New Participations"].max_row == 3


def test_changed_values_are_reported_and_repeat_is_empty(report_context):
    scraper.sync_event(
        "GIFA", 2023, "GIFA2023",
        [_record(website="https://old.example")],
    )
    scraper.sync_event(
        "GIFA", 2023, "GIFA2023",
        [_record(website="https://new.example", location="Hall 2 / B2")],
    )
    reports = sorted((report_context / "GIFA").glob("sync_changes_*.xlsx"))
    workbook = load_workbook(reports[-1], read_only=True)
    rows = list(workbook["Changed Data"].iter_rows(min_row=2, values_only=True))
    changes = {(row[4], row[5], row[6]) for row in rows}

    assert ("website", "https://old.example", "https://new.example") in changes
    assert ("hall", "Hall 1", "Hall 2") in changes
    count_before = len(reports)
    scraper.sync_event(
        "GIFA", 2023, "GIFA2023",
        [_record(website="https://new.example", location="Hall 2 / B2")],
    )
    assert len(list((report_context / "GIFA").glob("sync_changes_*.xlsx"))) == count_before


def test_failed_sync_does_not_retain_provisional_history_or_report(report_context):
    with pytest.raises(ValueError):
        scraper.sync_event("GIFA", 2023, "GIFA2023", [
            _record(),
            {"type": "profile"},
        ])

    with database.get_connection() as connection:
        assert connection.execute(
            "SELECT COUNT(*) FROM change_history"
        ).fetchone()[0] == 0
    assert not list(report_context.rglob("sync_changes_*.xlsx"))


def test_migration_creates_history_and_sales_fields_are_preserved(
    tmp_path
):
    database.DATABASE_PATH = tmp_path / "migration.db"
    database.DATA_DIR = tmp_path
    database.initialize_database()
    with database.get_connection() as connection:
        connection.execute(
            "INSERT INTO companies (name, normalized_name, created_at, updated_at, "
            "country, normalized_country, city, normalized_city, "
            "sales_status, sales_owner, sales_notes) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (
                "Acme GmbH", "acme gmbh", "now", "now", "Germany", "germany",
                "Berlin", "berlin", "hot", "owner", "note",
            ),
        )
        company_id = connection.execute(
            "SELECT id FROM companies"
        ).fetchone()[0]

    exhibition_id = database.get_or_create_exhibition(
        "GIFA", 2023, "GIFA2023"
    )
    database.sync_exhibitor(_record(), exhibition_id)

    with database.get_connection() as connection:
        row = connection.execute(
            "SELECT sales_status, sales_owner, sales_notes FROM companies "
            "WHERE id = ?", (company_id,)
        ).fetchone()
        assert connection.execute(
            "SELECT name FROM sqlite_master WHERE type='table' "
            "AND name='change_history'"
        ).fetchone() is not None
    assert tuple(row) == ("hot", "owner", "note")
